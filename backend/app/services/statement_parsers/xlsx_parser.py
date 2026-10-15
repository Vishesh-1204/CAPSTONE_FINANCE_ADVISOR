# app/services/statement_parsers/xlsx_parser.py
#
# Purpose: Excel (.xlsx) statements -> ParsedRows. openpyxl reads the first
# worksheet into a grid of cell VALUES (real datetime / number objects, not
# formatted text); tabular.assemble_rows does everything else.

import io
import logging
import zipfile

from openpyxl import load_workbook

from app.services.statement_parsers.base import ParseResult, StatementParser, StatementParseError
from app.services.statement_parsers.tabular import MAX_ROWS, assemble_rows

logger = logging.getLogger(__name__)


class XlsxParser(StatementParser):
    """Reads the first worksheet of an .xlsx file (the old binary .xls format is not supported)."""

    source_type = "xlsx"

    def parse(self, file_bytes: bytes, password: str | None = None) -> ParseResult:
        try:
            # read_only: streams rows (low memory); data_only: values instead of formulas
            workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
            logger.warning(f"Unreadable .xlsx upload ({type(exc).__name__})")
            raise StatementParseError("This is not a readable .xlsx file. Save it as .xlsx (Excel 2007+) or CSV.")

        try:
            sheet = workbook.worksheets[0]
            grid: list[list[object]] = []
            for values in sheet.iter_rows(values_only=True):
                grid.append(list(values))
                if len(grid) > MAX_ROWS + 100:             # header/preamble allowance
                    raise StatementParseError(f"Too many rows (limit {MAX_ROWS}). Split the file.")
        finally:
            workbook.close()

        if not grid:
            raise StatementParseError("The spreadsheet is empty.")
        return assemble_rows(grid, first_row_number=1)
